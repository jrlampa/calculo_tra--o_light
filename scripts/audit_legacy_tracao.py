import os
import shutil
import tempfile
import csv
import re
import sys
import openpyxl

# Add python directory to path to import ponto_blocks
sys.path.append(os.path.join(os.getcwd(), 'python'))
from translated.ponto_blocks import (
    calcular_polo,
    MTTraversalInput,
    BTTraversalInput,
    BTZeroTraversalInput,
    RamaisTraversalInput
)

TARGET_DIR = r"C:\Users\jonat\OneDrive - IM3 Brasil\LIGHT\PROJETOS\AP COSMO LINHA NOVA 3\CALC TRAÇÃO"
REPORT_FILE = "legacy_audit_report_ap_cosmo.csv"
TOL = 0.0001

def extract_resistance(modelo_str):
    """Extract numeric resistance from string like '11 m / 600 daN'."""
    if not modelo_str: return 0
    match = re.search(r'(\d+)\s*daN', str(modelo_str), re.IGNORECASE)
    if match:
        return float(match.group(1))
    return 0

def safe_float(v):
    if v is None: return 0.0
    if isinstance(v, (int, float)): return float(v)
    try: return float(str(v).replace(',', '.').strip())
    except: return 0.0

def audit_file(file_path):
    current_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    tmp_dir = os.path.join(current_dir, "tmp")
    if not os.path.exists(tmp_dir):
        os.makedirs(tmp_dir)
        
    basename = os.path.basename(file_path)
    temp_path = os.path.join(tmp_dir, f"audit_{basename}")
    
    try:
        shutil.copy2(file_path, temp_path)
        wb = openpyxl.load_workbook(temp_path, data_only=True)
        if "Ponto (1)" not in wb.sheetnames:
            wb.close()
            return {"status": "SKIPPED", "error": f"Sheet 'Ponto (1)' not found."}
            
        sheet = wb["Ponto (1)"]
        
        mapping = {
            "TIPO_POSTE": (10, 3), "MODELO": (11, 3), "H_POSTE": (17, 3), "H_ANC": (18, 3),
            "TOTAL_EXCEL": (140, 3), "MT1_ROW": 12, "MT2_ROW": 38, "BT_ROW": 64,
        }
        
        for r in range(1, 160):
            label = str(sheet.cell(r, 2).value or "").upper()
            if "MODELO" in label and "POSTE" in label: mapping["MODELO"] = (r, 3)
            if "TIPO" in label and "POSTE" in label: mapping["TIPO_POSTE"] = (r, 3)
            if "ALTURA" in label and "POSTE" in label and not mapping["H_POSTE"]: mapping["H_POSTE"] = (r, 3)
            if "ALTURA" in label and "ANCORAGEM" in label and not mapping["H_ANC"]: mapping["H_ANC"] = (r, 3)
            
            # Robust TOTAL search: Avoid 'VENTO' row
            if "TOTAL" in label and "daN" in label and "VENTO" not in label:
                if r > 100 or not mapping.get("TOTAL_EXCEL"): mapping["TOTAL_EXCEL"] = (r, 3)
            
            if "REDE" == label.strip():
                prev_label = str(sheet.cell(r-1, 2).value or "").upper()
                if "1º" in label or "1º" in prev_label: mapping["MT1_ROW"] = r
                if "2º" in label or "2º" in prev_label: mapping["MT2_ROW"] = r

        excel_tracao = safe_float(sheet.cell(*mapping["TOTAL_EXCEL"]).value)
        
        tipo_poste = sheet.cell(*mapping["TIPO_POSTE"]).value if mapping.get("TIPO_POSTE") else ""
        modelo_poste = sheet.cell(*mapping["MODELO"]).value
        resistencia_nominal = extract_resistance(modelo_poste)
        
        # MT1
        mt1_inputs = []
        r_base = mapping["MT1_ROW"]
        for col in [3, 6, 9, 12]:
            mt1_inputs.append(MTTraversalInput(
                tipo_rede=sheet.cell(r_base, col).value or "", tipo_cabo=sheet.cell(r_base+1, col).value or "",
                vao=safe_float(sheet.cell(r_base+2, col).value), flecha=safe_float(sheet.cell(r_base+3, col).value),
                angulo=safe_float(sheet.cell(r_base+4, col).value),
                altura_poste=safe_float(sheet.cell(mapping["H_POSTE"][0], 3).value),
                altura_ancoragem=safe_float(sheet.cell(mapping["H_ANC"][0], 3).value)
            ))

        # BT
        r_bt = mapping["BT_ROW"]
        for r in range(1, 160):
            l = str(sheet.cell(r, 2).value or "").upper()
            if "REDE" in l and r > 60: r_bt = r; break
        bt_inputs = [BTTraversalInput(altura_ancoragem=safe_float(sheet.cell(r_bt+6, 3).value))]
        for col in [6, 9, 12]:
             bt_inputs.append(BTTraversalInput(
                tipo_rede=sheet.cell(r_bt, col).value or "", tipo_cabo=sheet.cell(r_bt+1, col).value or "",
                vao=safe_float(sheet.cell(r_bt+2, col).value), flecha=safe_float(sheet.cell(r_bt+3, col).value),
                angulo=safe_float(sheet.cell(r_bt+4, col).value), altura_ancoragem=safe_float(sheet.cell(r_bt+6, col).value)
            ))
            
        # MT2
        r_mt2 = mapping["MT2_ROW"]
        for r in range(1, 160):
            l = str(sheet.cell(r, 2).value or "").upper()
            if "REDE" in l and 35 < r < 60: r_mt2 = r; break
        mt2_inputs = []
        for col in [3, 6, 9, 12]:
            mt2_inputs.append(MTTraversalInput(
                tipo_rede=sheet.cell(r_mt2, col).value or "", tipo_cabo=sheet.cell(r_mt2+1, col).value or "",
                vao=safe_float(sheet.cell(r_mt2+2, col).value), flecha=safe_float(sheet.cell(r_mt2+3, col).value),
                angulo=safe_float(sheet.cell(r_mt2+4, col).value),
                altura_poste=safe_float(sheet.cell(mapping["H_POSTE"][0], 3).value),
                altura_ancoragem=safe_float(sheet.cell(mapping["H_ANC"][0], 3).value)
            ))

        wb.close() 

        btz_inputs = [BTZeroTraversalInput() for _ in range(4)]
        ral_inputs = [RamaisTraversalInput() for _ in range(4)]
        
        try:
            py_res = calcular_polo(mt1_inputs, mt2_inputs, bt_inputs, btz_inputs, ral_inputs, tipo_poste=tipo_poste or "", modelo_poste=modelo_poste or "")
            py_tracao = py_res.total_tracao
        except Exception as e:
            return {"status": "ERROR", "error": f"Python Engine Error: {e}"}
            
        divergence = abs(py_tracao - excel_tracao)
        parity_status = "OK" if divergence <= TOL else "FAILED"
        
        human_errors = []
        # 🚩 HUMAN ERROR DETECTION: Structural Overload
        # Rule: Traction > Resistance * 1.05 (LIGHT business rule: 5% tolerance)
        if resistencia_nominal > 0 and excel_tracao > resistencia_nominal * 1.05:
            human_errors.append(f"SOBRECARGA: Esforco({excel_tracao:.2f}) > Resistencia({resistencia_nominal})")
            
        for i, inp in enumerate(mt1_inputs):
            if inp.vao > 0 and not inp.tipo_cabo: human_errors.append(f"MT1-T{i+1}: Vao s/ Cabo")
        for i, inp in enumerate(mt2_inputs):
            if inp.vao > 0 and not inp.tipo_cabo: human_errors.append(f"MT2-T{i+1}: Vao s/ Cabo")

        return {
            "status": parity_status, "human_errors": "; ".join(human_errors) if human_errors else "Nenhum",
            "py_tracao": py_tracao, "excel_tracao": excel_tracao, "divergence": divergence
        }
    finally:
        if os.path.exists(temp_path):
            try: os.remove(temp_path)
            except: pass

def main():
    print(f"--- INICIANDO AUDITORIA EM MASSA ---")
    print(f"Diretorio: {TARGET_DIR}")
    files = []
    for root, _, filenames in os.walk(TARGET_DIR):
        for f in filenames:
            if f.endswith(".xlsm") and not f.startswith("~$"):
                files.append(os.path.join(root, f))
    print(f"Arquivos encontrados: {len(files)}")
    with open(REPORT_FILE, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["Arquivo", "Status Paridade", "Erros Humanos Detectados", "Tracao Python", "Tracao Excel", "Divergencia"])
        ok_count = 0
        error_count = 0
        for i, file_path in enumerate(files):
            basename = os.path.basename(file_path)
            print(f"[{i+1}/{len(files)}] Auditando: {basename}...", end="\r")
            res = audit_file(file_path)
            if res.get("status") == "OK": ok_count += 1
            if res.get("human_errors") != "Nenhum" and res.get("status") != "SKIPPED": error_count += 1
            writer.writerow([basename, res.get("status", "ERROR"), res.get("human_errors", res.get("error", "Unknown")), f"{res.get('py_tracao', 0):.4f}", f"{res.get('excel_tracao', 0):.4f}", f"{res.get('divergence', 0):.6f}"])
    print(f"\n--- RESUMO DA AUDITORIA ---")
    print(f"Total Arquivos: {len(files)}")
    print(f"Paridade OK: {ok_count} ({ok_count/len(files)*100:.1f}%)")
    print(f"Arquivos com Erros Humanos: {error_count} ({error_count/len(files)*100:.1f}%)")
    print(f"Relatorio gerado: {REPORT_FILE}")

if __name__ == "__main__":
    main()
