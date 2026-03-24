
import openpyxl
import os
import json
from dataclasses import dataclass
from typing import List, Optional

@dataclass
class ExcelResult:
    filename: str
    projeto: str
    ponto: str
    tracao_total: float
    status_poste: str

def audit_file(filepath):
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True, keep_vba=True)
        sheet_names = wb.sheetnames
        target_sheet = next((s for s in sheet_names if s.lower() in ['ponto (1)', 'plan1']), None)
        
        if not target_sheet:
            print(f"Skipping {filepath}: Sheets {sheet_names} found.")
            return None
            
        ws = wb[target_sheet]
        print(f"Auditando aba: {target_sheet} em {os.path.basename(filepath)}")
        
        # Mapping for "Clandestino" layout
        def find_row(ws, keyword):
            for r in range(1, 400):
                val = str(ws.cell(row=r, column=2).value or "").upper()
                if keyword.upper() in val:
                    return r
            return None

        tr_row = find_row(ws, "Fração Total")
        st_row = find_row(ws, "Status Poste")
        
        ponto = ws.cell(row=1, column=11).value
        projeto = ws.cell(row=1, column=8).value
        tracao_total = ws.cell(row=tr_row, column=3).value if tr_row else 0.0
        status = ws.cell(row=st_row, column=3).value if st_row else "NOT FOUND"
        
        print(f"  Ponto: {ponto}, Projeto: {projeto}, Tração: {tracao_total}, Status: {status}")
        
        return ExcelResult(
            filename=os.path.basename(filepath),
            projeto=str(projeto),
            ponto=str(ponto),
            tracao_total=float(tracao_total) if tracao_total else 0.0,
            status_poste=str(status)
        )
    except Exception as e:
        print(f"Error auditing {filepath}: {e}")
        return None

def main():
    base_dir = r"C:\Users\jonat\OneDrive - IM3 Brasil\LIGHT\PROJETOS\REDE CLANDESTINA - RUAS JERUSALÉM E UVA - SANTA CRUZ RJ\CALC TRAÇÃO"
    folders = ["PROJETO I - RUA JERUSALÉM", "PROJETO II - RUA UVA"]
    
    audit_data = []
    
    for folder in folders:
        full_path = os.path.join(base_dir, folder)
        files = [f for f in os.listdir(full_path) if f.endswith('.xlsm')]
        
        # Pick first 3 from each
        for filename in files[:3]:
            res = audit_file(os.path.join(full_path, filename))
            if res:
                audit_data.append(res.__dict__)
                
    with open('parity_audit_reference.json', 'w', encoding='utf-8') as f:
        json.dump(audit_data, f, indent=4, ensure_ascii=False)
    
    print(f"Audit completed. {len(audit_data)} files processed.")

if __name__ == "__main__":
    main()
