import openpyxl
from openpyxl.utils import range_boundaries

def get_named_range_values(wb, name):
    try:
        if name not in wb.defined_names:
            return []
        nr = wb.defined_names[name]
        dests = list(nr.destinations)
        sheetname, coord = dests[0]
        ws = wb[sheetname]
        
        # se o range for algo como 'Plan4'!$A$2:$A$10
        cells = ws[coord.replace('$', '')]
        if hasattr(cells, '__iter__'):
             return [cell[0].value for cell in cells if cell[0].value]
        else:
             return [cells.value] if cells.value else []
    except Exception as e:
        return [f"Error: {str(e)}"]

wb = openpyxl.load_workbook('../AP COSMO LDA NOVA 03 - PROJETO 5 - POSTE 1D.xlsm', data_only=True)

print("Redes (Tipo):")
print(get_named_range_values(wb, 'Tipo'))

print("\nCabos (Tipo8):")
print(get_named_range_values(wb, 'Tipo8'))

print("\nPostes (Tipo9):")
print(get_named_range_values(wb, 'Tipo9'))

# Postes often have a dependent list. Let's look at Plan 4 directly.
if 'Plan4' in wb.sheetnames:
    ws4 = wb['Plan4']
    print("\nPlan4 - Coluna A (Redes?):")
    print([ws4.cell(row=i, column=1).value for i in range(1, 10)])
    print("\nPlan4 - Coluna F (Modelos Poste?):")
    print([ws4.cell(row=i, column=6).value for i in range(1, 15)])
