
import openpyxl
import os

files = [
    r"C:\Users\jonat\OneDrive - IM3 Brasil\LIGHT\PROJETOS\REDE CLANDESTINA - RUAS JERUSALÉM E UVA - SANTA CRUZ RJ\CALC TRAÇÃO\PROJETO I - RUA JERUSALÉM\CLANDESTINO - RUA JERUSALÉM - PROJETO I - POSTE 1.xlsm",
    r"C:\Users\jonat\OneDrive - IM3 Brasil\LIGHT\PROJETOS\REDE CLANDESTINA - RUAS JERUSALÉM E UVA - SANTA CRUZ RJ\CALC TRAÇÃO\PROJETO I - RUA JERUSALÉM\CLANDESTINO - RUA JERUSALÉM - PROJETO I - POSTE 15.xlsm",
    r"C:\Users\jonat\OneDrive - IM3 Brasil\LIGHT\PROJETOS\REDE CLANDESTINA - RUAS JERUSALÉM E UVA - SANTA CRUZ RJ\CALC TRAÇÃO\PROJETO II - RUA UVA\CLANDESTINO - RUA UVA - PROJETO II - POSTE 22.xlsm"
]

def find_val(ws, keyword):
    for r in range(1, 400):
        val = str(ws.cell(row=r, column=2).value or "").upper()
        if keyword.upper() in val:
            return ws.cell(row=r, column=3).value
    return None

for f in files:
    print(f"\nFILE: {os.path.basename(f)}")
    wb = openpyxl.load_workbook(f, data_only=True)
    ws = wb['Ponto (1)' if 'Ponto (1)' in wb.sheetnames else 'Plan1']
    print(f"  Ponto: {ws.cell(row=1, column=11).value}")
    print(f"  Tracao: {find_val(ws, 'Fração Total')}")
    print(f"  Status: {find_val(ws, 'Status Poste')}")
