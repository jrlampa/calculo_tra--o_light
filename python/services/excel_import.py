import logging
from io import BytesIO
from typing import Any, List
import openpyxl
from api.schemas import (
    CalculoInput,
    MTTraversalIn,
    BTTraversalIn,
    BTZeroTraversalIn,
    RamaisTraversalIn,
    PosteCalculoIn,
    CabecalhoIn
)

logger = logging.getLogger(__name__)

def safe_float(value: Any) -> float:
    """Safe conversion from Excel cell to float."""
    try:
        if value is None or str(value).strip() == "":
            return 0.0
        if isinstance(value, str):
            return float(value.replace(",", "."))
        return float(value)
    except (ValueError, TypeError):
        return 0.0

def find_row_by_keyword(ws, keyword: str, start_row: int = 1, col_range: List[int] = [1, 2]) -> int:
    """Find row index containing keyword in specified columns."""
    for row in range(start_row, min(ws.max_row, 350)):
        for c in col_range:
            val = str(ws.cell(row=row, column=c).value or "").strip().upper()
            # Be very strict with level headers to avoid picking up partial matches
            if keyword.upper() == "MT - 1" or keyword.upper() == "MT - 1º NÍVEL":
                if "MT - 1" in val:
                    return row
            elif keyword.upper() == "BT":
                # Strict BT to avoid BTZERO or RAMAIS BTZERO
                if val == "BT" or val == "BT:":
                    return row
            elif keyword.upper() in val:
                return row
    return -1

def extract_excel_to_input(file_content: bytes) -> CalculoInput:
    """Extracts engineering data from legacy XLSM/XLSX to CalculoInput schema."""
    wb = openpyxl.load_workbook(BytesIO(file_content), data_only=True)
    ws = wb.active

    # Find row anchors for levels using multi-column scan
    mt1_header = find_row_by_keyword(ws, "MT - 1", col_range=[1, 2, 3, 4])
    mt2_header = find_row_by_keyword(ws, "MT - 2", col_range=[1, 2, 3, 4])
    bt_header = find_row_by_keyword(ws, "BT", col_range=[1, 2, 3, 4])

    def find_sub_anchor_multi(ws, start_row, keyword, col_range=[1, 2, 3]):
        if start_row <= 0:
            return -1
        # Search up to 30 rows below header
        for r in range(start_row, start_row + 30):
             for c in col_range:
                 val = str(ws.cell(row=r, column=c).value or "").upper()
                 if keyword.upper() in val:
                     return r
        return -1

    def get_level_rows(header):
        if header <= 0:
            return None
        return {
            "rede": find_sub_anchor_multi(ws, header, "Tipo de rede"),
            "cabo": find_sub_anchor_multi(ws, header, "Tipo de cabo"),
            "vao": find_sub_anchor_multi(ws, header, "Vão"),
            "flecha": find_sub_anchor_multi(ws, header, "Flecha"),
            "angulo": find_sub_anchor_multi(ws, header, "Ângulo"),
            "h_poste": find_sub_anchor_multi(ws, header, "Altura poste") or find_sub_anchor_multi(ws, header, "Altura Poste"),
            "h_ancor": find_sub_anchor_multi(ws, header, "Altura ancoragem") or find_sub_anchor_multi(ws, header, "Altura Ancoragem")
        }

    rows_mt1 = get_level_rows(mt1_header)
    rows_mt2 = get_level_rows(mt2_header)
    rows_bt = get_level_rows(bt_header)

    # 1. Cabecalho logic improvement
    def find_field_val_flexible(ws, keyword, start_row=1, end_row=50):
        for r in range(start_row, end_row):
             for c in range(1, 8):
                  val = str(ws.cell(row=r, column=c).value or "").strip().upper()
                  if keyword.upper() in val:
                       res = ws.cell(row=r, column=c+1).value
                       if res is None:
                           res = ws.cell(row=r, column=c+2).value
                       return res
        return None

    projeto = find_field_val_flexible(ws, "Projeto:") or ws["C5"].value
    ponto = find_field_val_flexible(ws, "Ponto:") or ws["C6"].value
    endereco = find_field_val_flexible(ws, "Endereço:") or ws["C7"].value
    estudado_por = find_field_val_flexible(ws, "Estudado por:") or ws["C8"].value
    data_val = find_field_val_flexible(ws, "Data:") or ws["C10"].value

    cab_data = CabecalhoIn(
        projeto=str(projeto or ""), numero=str(ponto or ""),
        endereco=str(endereco or ""), estudado_por=str(estudado_por or ""), data=str(data_val or "")
    )

    # 2. Poste
    tipo_poste_val = find_field_val_flexible(ws, "Tipo do Poste") or ws["C140"].value
    modelo_poste_val = find_field_val_flexible(ws, "Modelo do Poste") or ws["B12"].value
    poste_data = PosteCalculoIn(
        tipo_poste=str(tipo_poste_val or ""),
        modelo_poste=str(modelo_poste_val or ""),
    )

    # 3. Traversals cols
    def get_traversal_cols_v3(ws, rows):
        if not rows or rows["rede"] <= 0:
            return [2, 5, 8, 11]
        label_col = -1
        for c in range(1, 8):
            if "TIPO DE REDE" in str(ws.cell(row=rows["rede"], column=c).value or "").upper():
                label_col = c
                break
        if label_col > 0:
            start_c = label_col + 1
            return [start_c, start_c + 3, start_c + 6, start_c + 9]
        return [2, 5, 8, 11]

    mt1_cols = get_traversal_cols_v3(ws, rows_mt1)
    mt2_cols = get_traversal_cols_v3(ws, rows_mt2)
    bt_cols  = get_traversal_cols_v3(ws, rows_bt)

    def get_mt_row_dyn(rows, cols, i):
        if not rows:
            return MTTraversalIn()
        col = cols[i]
        return MTTraversalIn(
            tipo_rede=str(ws.cell(row=rows["rede"], column=col).value or "") if rows["rede"] > 0 else "",
            tipo_cabo=str(ws.cell(row=rows["cabo"], column=col).value or "") if rows["cabo"] > 0 else "",
            vao=safe_float(ws.cell(row=rows["vao"], column=col).value) if rows["vao"] > 0 else 0.0,
            flecha=safe_float(ws.cell(row=rows["flecha"], column=col).value) if rows["flecha"] > 0 else 0.0,
            angulo=safe_float(ws.cell(row=rows["angulo"], column=col).value) if rows["angulo"] > 0 else 0.0,
            altura_poste=safe_float(ws.cell(row=rows["h_poste"], column=col).value) if rows["h_poste"] > 0 else 0.0,
            altura_ancoragem=safe_float(ws.cell(row=rows["h_ancor"], column=col).value) if rows["h_ancor"] > 0 else 0.0,
        )

    mt1 = [get_mt_row_dyn(rows_mt1, mt1_cols, i) for i in range(4)]
    mt2 = [get_mt_row_dyn(rows_mt2, mt2_cols, i) for i in range(4)]

    bt = []
    for i in range(4):
        if not rows_bt:
            bt.append(BTTraversalIn())
            continue
        col = bt_cols[i]
        bt.append(BTTraversalIn(
            tipo_rede=str(ws.cell(row=rows_bt["rede"], column=col).value or "") if rows_bt["rede"] > 0 else "",
            tipo_cabo=str(ws.cell(row=rows_bt["cabo"], column=col).value or "") if rows_bt["cabo"] > 0 else "",
            vao=safe_float(ws.cell(row=rows_bt["vao"], column=col).value) if rows_bt["vao"] > 0 else 0.0,
            flecha=safe_float(ws.cell(row=rows_bt["flecha"], column=col).value) if rows_bt["flecha"] > 0 else 0.0,
            angulo=safe_float(ws.cell(row=rows_bt["angulo"], column=col).value) if rows_bt["angulo"] > 0 else 0.0,
            altura_poste=safe_float(ws.cell(row=rows_bt["h_poste"], column=col).value) if rows_bt["h_poste"] > 0 else 0.0,
            altura_ancoragem=safe_float(ws.cell(row=rows_bt["h_ancor"], column=col).value) if rows_bt["h_ancor"] > 0 else 0.0,
        ))

    btz = [BTZeroTraversalIn() for _ in range(4)]
    ral = [RamaisTraversalIn() for _ in range(4)]

    return CalculoInput(cabecalho=cab_data, poste=poste_data, mt1=mt1, mt2=mt2, bt=bt, btz=btz, ral=ral)
