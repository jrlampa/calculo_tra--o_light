import openpyxl
import os

file_path = "AP COSMO LDA NOVA 03 - PROJETO 5 - POSTE 1D.xlsm"

if not os.path.exists(file_path):
    print(f"Arquivo não encontrado: {file_path}")
    exit(1)

wb = openpyxl.load_workbook(file_path, data_only=True)
sheet = wb["Ponto (1)"]

print(f"--- Data Validations in {file_path} ---")
for dv in sheet.data_validations.dataValidation:
    print(f"Cells: {dv.sqref}, Formula: {dv.formula1}")

# Common locations for dropdown lists in this template:
cells_to_check = ["B41", "C41", "B42", "C42"] # MT1 area
for c in cells_to_check:
    print(f"Value at {c}: {sheet[c].value}")

# Also check Plan4 or wherever the lists are stored.
if "Plan4" in wb.sheetnames:
    p4 = wb["Plan4"]
    print("\n--- Plan4 Content (Source of Lists) ---")
    for row in p4.iter_rows(min_row=1, max_row=40, min_col=1, max_col=10):
        # Only print non-empty rows
        vals = [cell.value for cell in row if cell.value is not None]
        if vals:
            print(vals)
